"""Unit tests for the story ingestion pipeline."""
import json
import os
import re
import pytest
from unittest.mock import patch, MagicMock
from pipeline.ingester import (
    extract_text_from_file,
    generate_header,
    redact_names,
    IngestState,
    StoryIngester,
)


# ── Text extraction ───────────────────────────────────────────────────────────

class TestExtractText:
    def test_extracts_plain_text(self, tmp_path):
        f = tmp_path / "story.txt"
        f.write_text("This is my story about leading a reorg.")
        text = extract_text_from_file(str(f))
        assert "leading a reorg" in text

    def test_extracts_markdown(self, tmp_path):
        f = tmp_path / "story.md"
        f.write_text("# My Story\n\nI led a team of 20 engineers.")
        text = extract_text_from_file(str(f))
        assert "led a team" in text

    def test_extracts_docx(self, tmp_path):
        from docx import Document
        doc = Document()
        doc.add_paragraph("Led a 40-person ads team through acquisition.")
        path = str(tmp_path / "story.docx")
        doc.save(path)
        text = extract_text_from_file(path)
        assert "40-person ads team" in text

    def test_unsupported_extension_raises(self, tmp_path):
        f = tmp_path / "story.pages"
        f.write_text("content")
        with pytest.raises(ValueError, match="Unsupported"):
            extract_text_from_file(str(f))

    def test_strips_leading_trailing_whitespace(self, tmp_path):
        f = tmp_path / "story.txt"
        f.write_text("\n\n  Some content  \n\n")
        text = extract_text_from_file(str(f))
        assert text == text.strip()


# ── Header generation ─────────────────────────────────────────────────────────

MOCK_HEADER_RESPONSE = """---
title: Reorganized 40-person ads eng org post-acquisition, reducing incidents 60% and growing revenue 35% YoY
type: story
keywords: [reorg, ads, team-growth, acquisition, managing-managers, platform, 2021]
summary: After acquiring AdCo in 2021, I inherited a fragmented 40-person ads org split across 3 locations with unclear ownership and two underperforming managers. I consolidated the org into 4 product-aligned teams over 6 months, replaced 2 managers, and established a new eng leadership cadence with weekly leads syncs and quarterly roadmap reviews. Ads revenue grew 35% YoY in the following year and eng-driven incidents dropped 60%. This story demonstrates org design ability, managing managers, and delivering business outcomes through structural change.
relevance_hint: Select when JD mentions post-M&A integration, org design, or managing eng managers at scale
---"""

class TestGenerateHeader:
    def test_returns_yaml_frontmatter(self, mock_provider):
        mock_provider.complete_fast.return_value = MOCK_HEADER_RESPONSE
        result = generate_header("Some story text", provider=mock_provider)
        assert result.startswith("---")
        assert "title:" in result
        assert "type:" in result
        assert "keywords:" in result
        assert "summary:" in result
        assert "relevance_hint:" in result

    def test_summary_field_present(self, mock_provider):
        mock_provider.complete_fast.return_value = MOCK_HEADER_RESPONSE
        result = generate_header("Some story text", provider=mock_provider)
        assert "summary:" in result
        summary_match = re.search(r"summary:\s*(.+)", result)
        assert summary_match and len(summary_match.group(1)) > 50

    def test_full_file_content_sent_to_provider(self, mock_provider):
        """Verify the full content is passed — no truncation."""
        long_content = "Story content. " * 500  # ~7500 chars
        mock_provider.complete_fast.return_value = MOCK_HEADER_RESPONSE
        generate_header(long_content, provider=mock_provider)
        call_content = mock_provider.complete_fast.call_args[0][0]
        assert long_content[:100] in call_content
        assert long_content[-100:] in call_content

    def test_falls_back_on_malformed_response(self, mock_provider):
        mock_provider.complete_fast.return_value = "I cannot generate a header."
        result = generate_header("Some story text", provider=mock_provider)
        assert "title:" in result
        assert "type: other" in result
        assert "summary:" in result


# ── Name redaction ───────────────────────────────────────────────────────────

MOCK_REDACTED = "I worked with [DIRECT REPORT] and [PEER MANAGER] to deliver the project. [VP] approved the roadmap."

class TestRedactNames:
    def test_returns_string(self, mock_provider):
        mock_provider.complete_fast.return_value = MOCK_REDACTED
        result = redact_names("I worked with John Smith and Jane Doe.", provider=mock_provider)
        assert isinstance(result, str)

    def test_names_replaced_with_placeholders(self, mock_provider):
        mock_provider.complete_fast.return_value = MOCK_REDACTED
        result = redact_names("I worked with John Smith and Jane Doe.", provider=mock_provider)
        assert "John" not in result
        assert "Jane" not in result
        assert "[" in result

    def test_placeholders_are_role_based_not_generic(self, mock_provider):
        mock_provider.complete_fast.return_value = MOCK_REDACTED
        result = redact_names("I worked with John and reported to Jane.", provider=mock_provider)
        assert any(p in result for p in [
            "[DIRECT REPORT]", "[PEER MANAGER]", "[VP]", "[MANAGER]",
            "[COLLEAGUE]", "[ENGINEER]", "[STAKEHOLDER]", "[CLIENT]"
        ])

    def test_falls_back_to_original_on_error(self, mock_provider):
        mock_provider.complete_fast.side_effect = Exception("API error")
        original = "Story with John Smith in it."
        result = redact_names(original, provider=mock_provider)
        assert result == original

    def test_redaction_applied_before_header_generation(self, tmp_path, mock_provider):
        """Verify redacted text (not original) is what gets saved to review/."""
        inbox = tmp_path / "inbox"
        review = tmp_path / "review"
        stories = tmp_path / "stories"
        for d in [inbox, review, stories]:
            d.mkdir()

        (inbox / "story.txt").write_text("John Smith helped me deliver the reorg.")
        ingester = StoryIngester(
            inbox_dir=str(inbox),
            review_dir=str(review),
            stories_dir=str(stories),
            profile_dir=str(tmp_path),
            provider=mock_provider,
        )

        with patch("pipeline.ingester.redact_names", return_value="[DIRECT REPORT] helped me deliver the reorg.") as mock_redact, \
             patch("pipeline.ingester.generate_header", return_value=MOCK_HEADER_RESPONSE):
            ingester.process_inbox()

        mock_redact.assert_called_once()
        review_content = open(str(review / "story.md")).read()
        assert "John" not in review_content
        assert "[DIRECT REPORT]" in review_content


# ── IngestState ───────────────────────────────────────────────────────────────

class TestIngestState:
    def test_new_state_is_empty(self, tmp_path):
        state = IngestState(str(tmp_path / ".ingest_state.json"))
        assert state.is_processed("file.txt") is False

    def test_mark_and_check_processed(self, tmp_path):
        state = IngestState(str(tmp_path / ".ingest_state.json"))
        state.mark_processed("story.docx", checksum="abc123")
        assert state.is_processed("story.docx") is True

    def test_state_persists_across_instances(self, tmp_path):
        path = str(tmp_path / ".ingest_state.json")
        state1 = IngestState(path)
        state1.mark_processed("story.docx", checksum="abc123")
        state2 = IngestState(path)
        assert state2.is_processed("story.docx") is True

    def test_changed_checksum_marks_as_unprocessed(self, tmp_path):
        state = IngestState(str(tmp_path / ".ingest_state.json"))
        state.mark_processed("story.docx", checksum="abc123")
        assert state.is_processed("story.docx", checksum="xyz999") is False

    def test_list_processed_returns_filenames(self, tmp_path):
        state = IngestState(str(tmp_path / ".ingest_state.json"))
        state.mark_processed("a.docx", checksum="1")
        state.mark_processed("b.docx", checksum="2")
        names = state.list_processed()
        assert "a.docx" in names
        assert "b.docx" in names


# ── StoryIngester ─────────────────────────────────────────────────────────────

@pytest.fixture
def ingester(tmp_path, mock_provider):
    inbox = tmp_path / "inbox"
    review = tmp_path / "review"
    stories = tmp_path / "stories"
    inbox.mkdir()
    review.mkdir()
    stories.mkdir()
    return StoryIngester(
        inbox_dir=str(inbox),
        review_dir=str(review),
        stories_dir=str(stories),
        profile_dir=str(tmp_path),
        provider=mock_provider,
    )


class TestStoryIngester:
    def test_processes_new_txt_file(self, ingester, tmp_path):
        (tmp_path / "inbox" / "story.txt").write_text("I led a platform reorg.")
        with patch("pipeline.ingester.generate_header", return_value=MOCK_HEADER_RESPONSE):
            result = ingester.process_inbox()
        assert result["processed"] == 1
        assert os.path.exists(os.path.join(str(tmp_path), "review", "story.md"))

    def test_skips_already_processed_files(self, ingester, tmp_path):
        f = tmp_path / "inbox" / "story.txt"
        f.write_text("I led a platform reorg.")
        with patch("pipeline.ingester.generate_header", return_value=MOCK_HEADER_RESPONSE):
            ingester.process_inbox()
            result = ingester.process_inbox()  # second run
        assert result["processed"] == 0
        assert result["skipped"] == 1

    def test_review_file_contains_header_and_content(self, ingester, tmp_path):
        (tmp_path / "inbox" / "story.txt").write_text("I led a platform reorg.")
        with patch("pipeline.ingester.generate_header", return_value=MOCK_HEADER_RESPONSE), \
             patch("pipeline.ingester.redact_names", side_effect=lambda text, provider: text):
            ingester.process_inbox()
        review_file = os.path.join(str(tmp_path), "review", "story.md")
        content = open(review_file).read()
        assert "---" in content
        assert "title:" in content
        assert "I led a platform reorg." in content

    def test_unsupported_files_are_skipped(self, ingester, tmp_path):
        (tmp_path / "inbox" / "story.pages").write_text("content")
        result = ingester.process_inbox()
        assert result["processed"] == 0
        assert result["skipped"] == 1

    def test_finalize_copies_review_to_stories(self, ingester, tmp_path):
        review_file = tmp_path / "review" / "story.md"
        review_file.write_text(MOCK_HEADER_RESPONSE + "\n\nStory content here.")
        ingester.finalize()
        assert os.path.exists(os.path.join(str(tmp_path), "stories", "story.md"))

    def test_finalize_builds_index(self, ingester, tmp_path):
        review_file = tmp_path / "review" / "story.md"
        review_file.write_text(MOCK_HEADER_RESPONSE + "\n\nStory content here.")
        ingester.finalize()
        index_path = os.path.join(str(tmp_path), "stories_index.md")
        assert os.path.exists(index_path)
        index = open(index_path).read()
        assert "reorg" in index.lower() or "ads" in index.lower()

    def test_finalize_overwrites_existing_stories(self, ingester, tmp_path):
        # Put old file in stories/
        old = tmp_path / "stories" / "story.md"
        old.write_text("old content")
        # Put updated file in review/
        review_file = tmp_path / "review" / "story.md"
        review_file.write_text(MOCK_HEADER_RESPONSE + "\n\nNew content.")
        ingester.finalize()
        assert "New content." in open(str(old)).read()

    def test_index_entry_format(self, ingester, tmp_path):
        review_file = tmp_path / "review" / "story.md"
        review_file.write_text(MOCK_HEADER_RESPONSE + "\n\nStory content here.")
        ingester.finalize()
        index = open(os.path.join(str(tmp_path), "stories_index.md")).read()
        assert "story.md" in index
        assert "story" in index.lower()  # type

    def test_index_includes_summary(self, ingester, tmp_path):
        review_file = tmp_path / "review" / "story.md"
        review_file.write_text(MOCK_HEADER_RESPONSE + "\n\nStory content here.")
        ingester.finalize()
        index = open(os.path.join(str(tmp_path), "stories_index.md")).read()
        # Summary content should appear in index
        assert "incidents dropped 60%" in index or "35% YoY" in index

    def test_process_inbox_returns_stats(self, ingester, tmp_path):
        (tmp_path / "inbox" / "a.txt").write_text("Story A")
        (tmp_path / "inbox" / "b.txt").write_text("Story B")
        (tmp_path / "inbox" / "c.pages").write_text("Unsupported")
        with patch("pipeline.ingester.generate_header", return_value=MOCK_HEADER_RESPONSE):
            result = ingester.process_inbox()
        assert result["processed"] == 2
        assert result["skipped"] == 1
        assert result["errors"] == 0
